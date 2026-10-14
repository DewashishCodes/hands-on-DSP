"""
report_kit.py
-------------
Shared Word-report builder for all DSP lab experiments.

Gives every experiment the same research-paper look:
  * Symbiosis Institute of Technology header on every page, page X of Y footer
  * numbered headings, captioned tables and figures (auto-numbered)
  * native, editable Word equations from LaTeX (via MathML -> OMML)
  * "Note" and "Inference" call-out boxes
  * syntax-highlighted code listings
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

ASSETS = Path(__file__).resolve().parent / "assets"
HEADER_IMAGE = ASSETS / "sit_header.jpeg"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"
DASHES = "[\u2013\u2014]"          # en dash, em dash: banned by house style
ACCENT = RGBColor(0x8B, 0x1A, 0x1A)       # Symbiosis maroon/red for headings
INK2 = RGBColor(0x40, 0x40, 0x40)
TEXT_W_CM = 17.0                          # A4 (21 cm) minus 2 x 2 cm margins

# Pygments token colours (VS-Code light-ish palette, print-safe)
CODE_COLOURS = {
    "Keyword": "0000C0", "Name.Builtin": "795E26", "Name.Function": "795E26",
    "Name.Class": "267F99", "Name.Decorator": "AF00DB", "String": "A31515",
    "Comment": "008000", "Number": "098658", "Operator": "000000", "Name": "001080",
}


# ==========================================================================
# Low-level XML helpers
# ==========================================================================
def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _cell_borders(cell, **edges):
    """edges: top/bottom/left/right = (size_eighths_pt, hex_colour) or None for nil."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        spec = edges.get(edge)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:color"), spec[1])
        borders.append(el)
    tc_pr.append(borders)


def _cell_margins(cell, top=60, bottom=60, left=110, right=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def _table_borders(table, size=4, colour="808080", inside=True):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge.startswith("inside") and not inside:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), colour)
        borders.append(el)
    tbl_pr.append(borders)


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _field(run, instr):
    """Insert a Word field (PAGE, NUMPAGES, ...) into a run."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(text)
    run._r.append(end)


def _keep_with_next(paragraph, on=True):
    paragraph.paragraph_format.keep_with_next = on


# ==========================================================================
# LaTeX -> native Word equation
# ==========================================================================
_xslt = None


def latex_to_omml(latex):
    global _xslt
    from latex2mathml.converter import convert
    if _xslt is None:
        _xslt = etree.XSLT(etree.parse(str(MML2OMML)))
    mathml = convert(latex)
    # Word's converter drops <mspace>, so \quad / \qquad would vanish: use real em-spaces
    mathml = re.sub(r'<mspace width="([0-9.]+)em"\s*/>',
                    lambda m: "<mtext>" + (" " if float(m.group(1)) < 0.5
                                           else " " * round(float(m.group(1)))) + "</mtext>", mathml)
    omml = _xslt(etree.fromstring(mathml))
    return omml.getroot()


# ==========================================================================
# Report builder
# ==========================================================================
class Report:
    def __init__(self, footer_left="DSP Laboratory"):
        self.doc = Document()
        self.fig_no = 0
        self.tab_no = 0
        self.eq_no = 0
        self.h = [0, 0, 0]
        self._setup_page(footer_left)
        self._setup_styles()

    # ---------------------------------------------------------------- setup
    def _setup_page(self, footer_left):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = Cm(4.55)
        s.bottom_margin = Cm(2.0)
        s.header_distance = Cm(0.7)
        s.footer_distance = Cm(0.9)

        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.add_run().add_picture(str(HEADER_IMAGE), width=Cm(TEXT_W_CM))

        fp = s.footer.paragraphs[0]
        tabs = fp.paragraph_format.tab_stops
        for pos in (Cm(8.255), Cm(16.51)):             # remove the Footer style's built-in tabs
            tabs.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)
        tabs.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
        r = fp.add_run(footer_left)
        r.font.size = Pt(8.5)
        r.font.color.rgb = INK2
        r = fp.add_run("\tPage ")
        r.font.size = Pt(8.5)
        r.font.color.rgb = INK2
        for instr, tail in (("PAGE", " of "), ("NUMPAGES", "")):
            r = fp.add_run()
            r.font.size = Pt(8.5)
            r.font.color.rgb = INK2
            _field(r, instr)
            if tail:
                t = fp.add_run(tail)
                t.font.size = Pt(8.5)
                t.font.color.rgb = INK2
        # thin rule above the footer
        p_pr = fp._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", "A0A0A0")):
            top.set(qn(f"w:{k}"), v)
        bdr.append(top)
        p_pr.append(bdr)

    def _setup_styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = BODY_FONT
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        normal.font.size = Pt(11)
        pf = normal.paragraph_format
        pf.space_after = Pt(5)
        pf.line_spacing = 1.12

        for lvl, size, before, after in ((1, 13.5, 14, 5), (2, 11.5, 10, 4), (3, 11, 8, 3)):
            hs = st[f"Heading {lvl}"]
            hs.font.name = BODY_FONT
            hs.element.rPr.rFonts.set(qn("w:asciiTheme"), "")  # drop theme font binding
            rfonts = hs.element.rPr.rFonts
            for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rfonts.get(qn(att)) is not None:
                    del rfonts.attrib[qn(att)]
            rfonts.set(qn("w:ascii"), BODY_FONT)
            rfonts.set(qn("w:hAnsi"), BODY_FONT)
            hs.font.size = Pt(size)
            hs.font.bold = True
            hs.font.italic = False
            hs.font.color.rgb = ACCENT if lvl == 1 else RGBColor(0x1F, 0x1F, 0x1F)
            hs.paragraph_format.space_before = Pt(before)
            hs.paragraph_format.space_after = Pt(after)
            hs.paragraph_format.keep_with_next = True

        cap = st["Caption"]
        cap.font.name = BODY_FONT
        cap.font.size = Pt(9.5)
        cap.font.italic = False
        cap.font.bold = False
        cap.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)

        for name in ("List Bullet", "List Number"):
            ls = st[name]
            ls.font.name = BODY_FONT
            ls.font.size = Pt(11)
            ls.paragraph_format.space_after = Pt(2.5)

    # ------------------------------------------------------------- inline text
    _TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\$[^$]+?\$)")

    def _runs(self, paragraph, text, size=None, colour=None):
        """Light markup: **bold**, *italic*, `code`, $inline LaTeX math$."""
        for part in self._TOKEN.split(text):
            if not part:
                continue
            if part.startswith("$") and part.endswith("$"):
                paragraph._p.append(latex_to_omml(part[1:-1]))
                continue
            if part.startswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("`"):
                run = paragraph.add_run(part[1:-1])
                run.font.name = CODE_FONT
                run.font.size = Pt((size or 11) - 1.5)
            elif part.startswith("*"):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            else:
                run = paragraph.add_run(part)
            if size and not part.startswith("`"):
                run.font.size = Pt(size)
            if colour:
                run.font.color.rgb = colour
        return paragraph

    # --------------------------------------------------------------- blocks
    def heading(self, text, level=1, numbered=True):
        if numbered:
            self.h[level - 1] += 1
            for i in range(level, 3):
                self.h[i] = 0
            num = ".".join(str(x) for x in self.h[:level])
            text = f"{num}  {text}" if level > 1 else f"{num}.  {text}"
        return self.doc.add_heading(text, level=level)

    def para(self, text, align="justify", size=None, space_after=None, keep_next=False, colour=None):
        p = self.doc.add_paragraph()
        p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        self._runs(p, text, size, colour)
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        _keep_with_next(p, keep_next)
        return p

    def bullets(self, items, numbered=False):
        style = "List Number" if numbered else "List Bullet"
        for it in items:
            p = self.doc.add_paragraph(style=style)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self._runs(p, it)

    def equation(self, latex, label=None, numbered=True):
        """Centred display equation with its number right-aligned: (n)."""
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.tab_stops.add_tab_stop(Cm(TEXT_W_CM / 2), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
        pf.space_before = Pt(3)
        pf.space_after = Pt(6)
        p.add_run("\t")
        omml = latex_to_omml(latex)
        for lim in omml.iter(qn("m:limLoc")):      # display style: limits above/below Σ
            lim.set(qn("m:val"), "undOvr")
        p._p.append(omml)
        if numbered:
            self.eq_no += 1
            p.add_run(f"\t({self.eq_no})")
        if label:
            self.labels = getattr(self, "labels", {})
            self.labels[label] = self.eq_no
        return self.eq_no

    def callout(self, title, text, fill="F3F2EE", bar="8B1A1A", items=None):
        """Single-cell shaded box with a coloured left bar. ``text`` may be a list of paragraphs."""
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.width = Cm(TEXT_W_CM)
        _shade(cell, fill)
        _cell_borders(cell, left=(24, bar), top=None, bottom=None, right=None)
        _cell_margins(cell, 90, 90, 160, 140)
        _no_split(t.rows[0])
        paras = text if isinstance(text, list) else [text]
        first = cell.paragraphs[0]
        r = first.add_run(f"{title}  ")
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(bar)
        first.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self._runs(first, paras[0], size=10)
        first.paragraph_format.space_after = Pt(2)
        for extra in paras[1:]:
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(2)
            self._runs(p, extra, size=10)
        for it in items or []:
            p = cell.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            self._runs(p, it, size=10)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def note(self, text, title="Note."):
        return self.callout(title, text, fill="EEF4FB", bar="2A78D6")

    def inference(self, text, items=None):
        return self.callout("Inference.", text, fill="EEF7F1", bar="1B8F5F", items=items)

    def caption(self, kind, title, desc=None):
        if kind == "Figure":
            self.fig_no += 1
            num = self.fig_no
        else:
            self.tab_no += 1
            num = self.tab_no
        p = self.doc.add_paragraph(style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind == "Table" else WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(f"{kind} {num}. ")
        r.bold = True
        r = p.add_run(title)
        r.bold = True
        if desc:
            p.add_run("  ")
            self._runs(p, desc, size=9.5)
        p.paragraph_format.space_after = Pt(8 if kind == "Figure" else 3)
        p.paragraph_format.space_before = Pt(2 if kind == "Figure" else 6)
        if kind == "Table":
            _keep_with_next(p)
        return num

    def figure(self, path, title, desc=None, width_cm=TEXT_W_CM, inference=None, inference_items=None):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        _keep_with_next(p)
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        num = self.caption("Figure", title, desc)
        if inference or inference_items:
            _keep_with_next(self.doc.paragraphs[-1])     # keep the inference with its figure
            self.inference(inference or "", inference_items)
        return num

    def table(self, header, rows, title, widths=None, desc=None, font_size=9.5, align=None,
              first_col_bold=False, zebra=True):
        """Booktabs-like table: shaded header row, thin rules, optional zebra rows."""
        self.caption("Table", title, desc)
        ncol = len(header)
        widths = widths or [TEXT_W_CM / ncol] * ncol
        t = self.doc.add_table(rows=1 + len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _table_borders(t, size=4, colour="A6A6A6")
        align = align or (["left"] + ["center"] * (ncol - 1))
        amap = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT}
        for r_i, values in enumerate([header] + list(rows)):
            row = t.rows[r_i]
            _no_split(row)
            if r_i == 0:
                _repeat_header(row)
            for c_i, val in enumerate(values):
                cell = row.cells[c_i]
                cell.width = Cm(widths[c_i])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _cell_margins(cell, 40, 40, 90, 90)
                p = cell.paragraphs[0]
                p.alignment = amap[align[c_i]] if r_i else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if r_i == 0:
                    _shade(cell, "8B1A1A")
                    run = p.add_run(str(val))
                    run.bold = True
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    if zebra and r_i % 2 == 0:
                        _shade(cell, "F5F4F0")
                    self._runs(p, str(val), size=font_size)
                    if first_col_bold and c_i == 0:
                        for run in p.runs:
                            run.bold = True
        if len(rows) <= 15:                    # short table: keep it on one page
            for row in t.rows[:-1]:
                for cell in row.cells:
                    for par in cell.paragraphs:
                        _keep_with_next(par)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def info_block(self, pairs):
        """Two-column key/value block (student details) without visible grid."""
        t = self.doc.add_table(rows=(len(pairs) + 1) // 2, cols=4)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(t, size=4, colour="BFBFBF", inside=False)
        widths = [3.2, 5.3, 3.2, 5.3]
        for i, (k, v) in enumerate(pairs):
            row = t.rows[i // 2]
            kc, vc = row.cells[(i % 2) * 2], row.cells[(i % 2) * 2 + 1]
            for c, w in ((kc, widths[(i % 2) * 2]), (vc, widths[(i % 2) * 2 + 1])):
                c.width = Cm(w)
                _cell_margins(c, 50, 50, 100, 80)
            kc.paragraphs[0].add_run(k).bold = True
            vc.paragraphs[0].add_run(v)
            for c in (kc, vc):
                c.paragraphs[0].paragraph_format.space_after = Pt(0)
                c.paragraphs[0].runs[0].font.size = Pt(10.5)
        return t

    def title(self, text, subtitle=None):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = ACCENT
        if subtitle:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            r = p.add_run(subtitle)
            r.italic = True
            r.font.size = Pt(11.5)
            r.font.color.rgb = INK2

    def abstract(self, text, keywords):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.width = Cm(TEXT_W_CM)
        _shade(cell, "FAFAF8")
        _cell_borders(cell, top=(8, "8B1A1A"), bottom=(8, "8B1A1A"), left=None, right=None)
        _cell_margins(cell, 110, 110, 160, 160)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run("Abstract: ")
        r.bold = True
        r.italic = True
        r.font.size = Pt(10)
        self._runs(p, text, size=10)
        p.paragraph_format.space_after = Pt(4)
        k = cell.add_paragraph()
        r = k.add_run("Keywords: ")
        r.bold = True
        r.italic = True
        r.font.size = Pt(10)
        r = k.add_run(keywords)
        r.italic = True
        r.font.size = Pt(10)
        k.paragraph_format.space_after = Pt(0)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def code(self, source, title=None, font_size=7.5, line_numbers=True):
        """Syntax-highlighted listing in a shaded box (breaks across pages)."""
        from pygments.lexers import PythonLexer
        from pygments.token import Token

        if title:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            r = p.add_run(title)
            r.bold = True
            r.font.size = Pt(9.5)
            p.paragraph_format.space_after = Pt(2)
            _keep_with_next(p)

        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.width = Cm(TEXT_W_CM)
        _shade(cell, "F7F7F5")
        _cell_borders(cell, top=(4, "C8C8C8"), bottom=(4, "C8C8C8"), left=(4, "C8C8C8"),
                      right=(4, "C8C8C8"))
        _cell_margins(cell, 70, 70, 120, 120)

        def colour_for(tok):
            while tok is not Token:
                key = str(tok).replace("Token.", "")
                if key in CODE_COLOURS:
                    return CODE_COLOURS[key]
                if key.startswith("Literal.String") or key.startswith("String"):
                    return CODE_COLOURS["String"]
                if key.startswith("Literal.Number"):
                    return CODE_COLOURS["Number"]
                if key.startswith("Comment"):
                    return CODE_COLOURS["Comment"]
                if key.startswith("Keyword"):
                    return CODE_COLOURS["Keyword"]
                tok = tok.parent
            return None

        lines = [[]]
        for tok, val in PythonLexer().get_tokens(source.rstrip("\n")):
            parts = val.split("\n")
            for i, part in enumerate(parts):
                if i:
                    lines.append([])
                if part:
                    lines[-1].append((tok, part))
        width = len(str(len(lines)))
        keep = len(lines) <= 40                 # short listing: keep it on one page
        first = True
        for ln, toks in enumerate(lines, 1):
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            pf = p.paragraph_format
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            pf.line_spacing = 1.0
            pf.keep_with_next = keep and ln < len(lines)
            if line_numbers:
                r = p.add_run(f"{ln:>{width}}  ")
                r.font.name = CODE_FONT
                r.font.size = Pt(font_size)
                r.font.color.rgb = RGBColor(0xA0, 0xA0, 0xA0)
            for tok, val in toks:
                r = p.add_run(val)
                r.font.name = CODE_FONT
                r.font.size = Pt(font_size)
                col = colour_for(tok)
                if col:
                    r.font.color.rgb = RGBColor.from_string(col)
                if "Comment" in str(tok):
                    r.italic = True
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def spacer(self, pt=4):
        self.doc.add_paragraph().paragraph_format.space_after = Pt(pt)

    def save(self, path):
        # House rule: no em dashes or en dashes anywhere in the document
        xml = "".join(part.blob.decode("utf-8", "ignore") for part in self.doc.part.package.parts
                      if part.partname.endswith(".xml"))
        bad = [m.start() for m in re.finditer(DASHES, xml)]
        if bad:
            snippet = re.sub("<[^>]+>", "", xml[max(0, bad[0] - 300): bad[0] + 60])[-90:]
            raise ValueError(f"{len(bad)} em/en dash(es) in the document, first near: ...{snippet}")
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path


def export_pdf(docx_path, pdf_path=None):
    """Convert with Microsoft Word (fields such as page numbers are refreshed)."""
    import subprocess
    docx_path = Path(docx_path).resolve()
    pdf_path = Path(pdf_path or docx_path.with_suffix(".pdf")).resolve()
    script = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{docx_path}'); $d.Fields.Update() | Out-Null; "
        f"$d.SaveAs2('{pdf_path}', 17); $d.Close($false); $w.Quit()"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)
    return pdf_path
